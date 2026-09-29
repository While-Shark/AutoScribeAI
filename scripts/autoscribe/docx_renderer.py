"""Export a validated manual as a readable Word document."""
import tempfile
from pathlib import Path

from .validation import ValidationError, asset_path, read_json, validate_manual


def _page_number(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()._r
    for kind, text in (('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)):
        element = OxmlElement('w:fldChar' if kind else 'w:instrText' if text == ' PAGE ' else 'w:t')
        if kind:
            element.set(qn('w:fldCharType'), kind)
        elif text == ' PAGE ':
            element.set(qn('xml:space'), 'preserve')
        if text:
            element.text = text
        run.append(element)


def _add_metadata(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = 3
    run = paragraph.add_run(label)
    run.bold = True
    paragraph.add_run(value)


def build_docx(manual, asset_root, output_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal']._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', 'Microsoft YaHei')
    styles['Normal'].paragraph_format.space_after = Pt(6)
    for style_name, size in (('Title', 25), ('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
        style = styles[style_name]
        style.font.name = 'Aptos Display' if style_name != 'Heading 3' else 'Aptos'
        style.font.size = Pt(size)
        style.font.bold = style_name != 'Title'
        style.font.color.rgb = RGBColor(0x18, 0x26, 0x3A)
        style._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', 'Microsoft YaHei')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('AutoScribeAI · ')
    _page_number(footer)
    doc.core_properties.title = manual['title']
    doc.core_properties.subject = f"{manual['project']['name']} 操作手册"

    title = doc.add_paragraph(style='Title')
    title.add_run(manual['title'])
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    intro.add_run(f"{manual['project']['name']} · {manual['project'].get('environment', '')}").bold = True
    _add_metadata(doc, '项目版本：', manual['project'].get('version', '未注明'))
    _add_metadata(doc, '适用角色：', '、'.join(manual['roles']))
    counts = {key: sum(item['status'] == key for item in manual['workflows'])
              for key in ('verified', 'blocked', 'unverified', 'pending')}
    _add_metadata(doc, '流程状态：', f"已验证 {counts['verified']} · 受阻 {counts['blocked']} · 未验证 {counts['unverified']} · 待处理 {counts['pending']}")
    doc.add_paragraph('本手册保留未验证和受阻流程的状态与原因；候选流程不会被当作真实操作结果。')
    doc.add_page_break()

    modules = {item['id']: item for item in manual['modules']}
    features_by_module = {ident: [] for ident in modules}
    for feature in manual['features']:
        features_by_module[feature['moduleId']].append(feature)
    workflows_by_feature = {item['id']: [] for item in manual['features']}
    for workflow in manual['workflows']:
        workflows_by_feature[workflow['featureId']].append(workflow)
    steps_by_workflow = {item['id']: [] for item in manual['workflows']}
    for step in manual['steps']:
        steps_by_workflow[step['workflowId']].append(step)
    evidence = {item['id']: item for item in manual['evidence']}
    chapters = {item['moduleId']: item for item in manual['chapters']}
    labels = {'verified': '已验证', 'blocked': '受阻', 'unverified': '未验证', 'pending': '待处理'}
    first_module = True
    with tempfile.TemporaryDirectory(prefix='autoscribe-docx-assets-') as working:
        working = Path(working)
        for module in manual['modules']:
            if not first_module:
                doc.add_page_break()
            first_module = False
            chapter = chapters.get(module['id'], {})
            doc.add_heading(chapter.get('title', module['name']), level=1)
            if chapter.get('purpose'):
                doc.add_paragraph(chapter['purpose'])
            _add_metadata(doc, '入口：', module['location'])
            _add_metadata(doc, '来源：', module['source'])
            for feature in features_by_module[module['id']]:
                doc.add_heading(feature['name'], level=2)
                _add_metadata(doc, '功能入口：', feature['location'])
                for workflow in workflows_by_feature[feature['id']]:
                    doc.add_heading(workflow['goal'], level=3)
                    _add_metadata(doc, '状态：', labels[workflow['status']])
                    _add_metadata(doc, '适用角色：', workflow['role'])
                    if workflow['location']:
                        _add_metadata(doc, '流程位置：', workflow['location'])
                    _add_metadata(doc, '成功标准：', workflow['successCriteria'])
                    if workflow['preconditions']:
                        doc.add_paragraph('前置条件', style='Heading 4')
                        for condition in workflow['preconditions']:
                            doc.add_paragraph(condition, style='List Bullet')
                    if workflow['status'] in ('blocked', 'unverified'):
                        paragraph = doc.add_paragraph()
                        paragraph.add_run('说明：').bold = True
                        paragraph.add_run(workflow.get('reason', '尚未验证'))
                    steps = sorted(steps_by_workflow[workflow['id']], key=lambda item: item['order'])
                    if not steps:
                        doc.add_paragraph('尚无实际操作步骤。源码或说明中的候选流程不能替代真实界面验证。')
                    for step in steps:
                        paragraph = doc.add_paragraph()
                        paragraph.paragraph_format.keep_with_next = True
                        paragraph.add_run(f"{step['order']}. ")
                        paragraph.add_run(step['action']).bold = True
                        _add_metadata(doc, '操作位置：', step['location'])
                        _add_metadata(doc, '预期结果：', step['expectedResult'])
                        if step.get('actualResult'):
                            _add_metadata(doc, '实际结果：', step['actualResult'])
                        for ident in step['evidenceIds']:
                            item = evidence[ident]
                            source = asset_path(asset_root, item['path'])
                            image_path = working / f"{ident}{source.suffix.lower()}"
                            image_path.write_bytes(source.read_bytes())
                            paragraph = doc.add_paragraph()
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            paragraph.paragraph_format.keep_with_next = True
                            run = paragraph.add_run()
                            from PIL import Image
                            with Image.open(image_path) as screenshot:
                                width_px, height_px = screenshot.size
                            if height_px / width_px > 7.2 / 6.25:
                                picture = run.add_picture(str(image_path), height=Inches(7.2))
                            else:
                                picture = run.add_picture(str(image_path), width=Inches(6.25))
                            picture._inline.docPr.set('descr', f"{item['page']} | {ident}")
                            caption = doc.add_paragraph(f"图 {step['order']}  {item['page']}")
                            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            caption.paragraph_format.keep_together = True
                            caption.paragraph_format.space_after = Pt(10)
        doc.add_heading('覆盖范围与限制', level=1)
        if manual['limitations']:
            for limitation in manual['limitations']:
                doc.add_paragraph(limitation, style='List Bullet')
        else:
            doc.add_paragraph('没有额外限制说明。')
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)


def render_docx(manual_path, out_docx):
    manual_path, out_docx = Path(manual_path), Path(out_docx)
    manual = validate_manual(read_json(manual_path), manual_path.parent)
    if out_docx.exists() or out_docx.is_symlink():
        raise ValidationError('DOCX 输出文件已存在；请使用新的路径，避免覆盖用户文件')
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.autoscribe-docx-', dir=out_docx.parent) as directory:
        staged = Path(directory) / 'manual.docx'
        build_docx(manual, manual_path.parent, staged)
        staged.replace(out_docx)
    return {'path': out_docx.name, 'evidenceCount': len(manual['evidence'])}
