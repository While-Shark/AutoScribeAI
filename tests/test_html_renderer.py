import base64
import hashlib
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from autoscribe.html_renderer import render_html
from autoscribe.standalone_html import export_standalone_html
from autoscribe.inventory import create_coverage_plan, coverage_report, inventory_to_manual
from autoscribe.state import atomic_json
from autoscribe.validation import ValidationError, read_json


class Links(HTMLParser):
    def __init__(self):
        super().__init__(); self.links=[]; self.images=[]; self.scripts=[]
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag in ('a','img'):
            (self.links if tag=='a' else self.images).append(attrs)
        if tag in ('script','link'):
            self.scripts.append((tag,attrs))


class HtmlRendererTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.config=read_json(ROOT/'examples/project.json')
        self.config['language'] = 'zh-CN'  # Keep baseline assertions Chinese; language matrix below covers English/default.
        self.config_path=self.root/'project.json'; self.config_path.write_text(json.dumps(self.config))
        self.inventory=read_json(ROOT/'examples/inventory.json')
        self.inventory_path=self.root/'inventory.json'; self.inventory_path.write_text(json.dumps(self.inventory))
        self.manual_path=self.root/'manual.json'; self.coverage_path=self.root/'coverage.json'
        self.plan_path=self.root/'plan.json'; self.out=self.root/'html'

    def prepare(self, verified=False):
        manual=inventory_to_manual(self.inventory,self.config)
        if verified:
            flow=manual['workflows'][0]; flow['status']='verified'; flow['stepIds']=['step-1']
            manual['chapters'][0]['faqs']=[{
                'question':'怎样打开列表？ <script>bad()</script>',
                'answer':'从左侧菜单选择“列表”。',
                'source':'observed', 'workflowIds':[flow['id']],
            }]
            (self.root/'evidence').mkdir(exist_ok=True)
            image_path=self.root/'evidence'/'screen.png'
            Image.new('RGB',(96,64),(235,240,250)).save(image_path)
            image_bytes=image_path.read_bytes()
            manual['steps']=[{'id':'step-1','workflowId':flow['id'],'order':1,'action':'打开列表',
                'location':'左侧导航','expectedResult':'显示条目','actualResult':'列表已显示',
                'source':'observed','evidenceIds':['ev-1']}]
            flow['location']='导航/列表'
            manual['evidence']=[{'id':'ev-1','stepIds':['step-1'],'path':'evidence/screen.png',
                'capturedAt':'2026-09-29T08:00:00Z','page':'条目列表',
                'viewport':{'width':96,'height':64},'source':'observed',
                'sha256':hashlib.sha256(image_bytes).hexdigest(),'redacted':True}]
        atomic_json(self.manual_path,manual)
        create_coverage_plan(self.inventory_path,self.config_path,self.plan_path)
        coverage_report(self.plan_path,self.inventory_path,self.config_path,self.manual_path,self.coverage_path)

    def test_single_file_html_keeps_evidence_after_moving_without_assets(self):
        self.prepare(verified=True)
        image_bytes = (self.root/'evidence/screen.png').read_bytes()
        portable = self.root/'portable.html'
        result = export_standalone_html(self.manual_path,self.coverage_path,portable)
        self.assertEqual(result['evidenceCount'],1)
        self.assertEqual(result['bytes'],portable.stat().st_size)
        with self.assertRaises(ValidationError):
            export_standalone_html(self.manual_path,self.coverage_path,portable)
        moved = self.root/'moved'/'portable.html'
        moved.parent.mkdir()
        portable.rename(moved)
        shutil.rmtree(self.root/'evidence')
        page = moved.read_text(encoding='utf-8')
        parser = Links(); parser.feed(page)
        embedded = [image['src'] for image in parser.images if image.get('src')]
        self.assertEqual(len(embedded),1)
        self.assertEqual(base64.b64decode(embedded[0].split(',',1)[1]),image_bytes)
        self.assertTrue(all(link['href'].startswith('#') for link in parser.links))
        self.assertNotIn('manual.docx',page)
        self.assertNotIn('coverage.json',page)
        self.assertIn('怎样打开列表？ &lt;script&gt;bad()&lt;/script&gt;',page)

    def test_offline_render_and_escaping_for_unverified_manual(self):
        self.prepare()
        manual=read_json(self.manual_path)
        manual['title']='</title><script>alert(1)</script>'
        atomic_json(self.manual_path,manual)
        coverage_report(self.plan_path,self.inventory_path,self.config_path,self.manual_path,self.coverage_path)
        report=render_html(self.manual_path,self.coverage_path,self.out)
        page=(self.out/'index.html').read_text()
        self.assertIn('&lt;/title&gt;&lt;script&gt;alert(1)&lt;/script&gt;',page)
        self.assertNotIn('</title><script>alert(1)',page)
        self.assertIn('离线手册',page)
        self.assertIn('0/3',page)
        self.assertEqual(report['ready'],False)
        self.assertEqual(report['copiedEvidenceCount'],0)
        parser=Links(); parser.feed(page)
        self.assertFalse(any(attrs.get('src','').startswith('http') for _,attrs in parser.scripts))

    def test_real_evidence_copied_with_alt_and_local_relative_path(self):
        self.prepare(verified=True)
        report=render_html(self.manual_path,self.coverage_path,self.out)
        self.assertEqual(report['verified'],1)
        self.assertEqual(report['copiedEvidenceCount'],1)
        self.assertTrue((self.out/'evidence/ev-1.png').is_file())
        parser=Links(); parser.feed((self.out/'index.html').read_text())
        self.assertTrue(any(i.get('src')=='evidence/ev-1.png' and i.get('alt') for i in parser.images))
        self.assertIn('showModal()', (self.out/'index.html').read_text())
        self.assertEqual(read_json(self.out/'coverage.json')['coverageDisplay'],'1/3')
        self.assertTrue((self.out/'manual.json').is_file())
        self.assertTrue((self.out/'quality-report.json').is_file())
        self.assertTrue((self.out/'manual.docx').is_file())
        self.assertTrue((self.out/'manual-markdown.zip').is_file())
        page=(self.out/'index.html').read_text()
        self.assertIn('href="manual.docx" download',page)
        self.assertIn('href="manual-markdown.zip" download',page)
        self.assertIn('常见问题',page)
        self.assertIn('怎样打开列表？ &lt;script&gt;bad()&lt;/script&gt;',page)
        self.assertNotIn('怎样打开列表？ <script>bad()',page)
        from docx import Document
        doc=Document(self.out/'manual.docx')
        self.assertTrue(any('打开列表' in paragraph.text for paragraph in doc.paragraphs))
        self.assertTrue(any('怎样打开列表？' in paragraph.text for paragraph in doc.paragraphs))
        doc_text='\n'.join(paragraph.text for paragraph in doc.paragraphs)
        markdown=''
        with zipfile.ZipFile(self.out/'manual-markdown.zip') as archive:
            self.assertEqual(set(archive.namelist()),{'README.md','assets/ev-1.png'})
            markdown=archive.read('README.md').decode('utf-8')
            self.assertIn('![步骤 1 · 条目列表](assets/ev-1.png)',markdown)
            self.assertIn('已验证',markdown)
            self.assertIn('怎样打开列表？ <script>bad()</script>',markdown)
            self.assertIn('答：从左侧菜单选择“列表”。',markdown)
            self.assertEqual(archive.read('assets/ev-1.png'),(self.root/'evidence/screen.png').read_bytes())
        for module in read_json(self.manual_path)['modules']:
            self.assertIn(module['name'],doc_text)
            self.assertIn(module['name'],markdown)
        workflow=read_json(self.manual_path)['workflows'][0]
        self.assertIn(workflow['goal'],doc_text)
        self.assertIn(workflow['goal'],markdown)
        self.assertEqual(len(doc.inline_shapes),1)

    def test_stale_coverage_rejected(self):
        self.prepare()
        manual=read_json(self.manual_path); manual['title']='Edited'; atomic_json(self.manual_path,manual)
        with self.assertRaisesRegex(ValidationError,'不匹配'):
            render_html(self.manual_path,self.coverage_path,self.out)
        self.assertFalse(self.out.exists())

    def test_existing_output_is_not_overwritten(self):
        self.prepare(); self.out.mkdir(); sentinel=self.out/'keep.txt'; sentinel.write_text('keep')
        with self.assertRaisesRegex(ValidationError,'已存在'):
            render_html(self.manual_path,self.coverage_path,self.out)
        self.assertEqual(sentinel.read_text(),'keep')

    def test_all_local_image_links_resolve(self):
        self.prepare(verified=True); render_html(self.manual_path,self.coverage_path,self.out)
        parser=Links(); parser.feed((self.out/'index.html').read_text())
        for image in parser.images:
            if not image.get('src'):
                continue
            self.assertTrue((self.out/image['src']).is_file())

    def test_exporters_localize_ui_for_supported_languages(self):
        cases = {
            'zh-CN': ('zh-CN', '目录', '常见问题', '已验证'),
            'en-US': ('en-US', 'Contents', 'Frequently asked questions', 'Verified'),
            'ja-JP': ('ja-JP', '目次', 'よくある質問', '検証済み'),
            'ko-KR': ('ko-KR', '목차', '자주 묻는 질문', '검증됨'),
        }
        for language, (html_lang, toc, faqs, verified) in cases.items():
            with self.subTest(language=language):
                self.config['language'] = language
                self.config_path.write_text(json.dumps(self.config), encoding='utf-8')
                self.out = self.root / f'html-{language}'
                self.prepare(verified=True)
                manual = read_json(self.manual_path)
                self.assertEqual(manual['project']['language'], language)
                if language == 'ja-JP':
                    self.assertIn('まだ実行・検証されていません', manual['workflows'][1]['reason'])
                if language == 'ko-KR':
                    self.assertIn('아직 실행 및 검증되지 않았습니다', manual['workflows'][1]['reason'])
                render_html(self.manual_path, self.coverage_path, self.out)

                page = (self.out / 'index.html').read_text(encoding='utf-8')
                self.assertIn(f'<html lang="{html_lang}">', page)
                self.assertIn(toc, page)
                self.assertIn(faqs, page)
                self.assertIn(verified, page)

                from docx import Document
                docx_text = '\n'.join(p.text for p in Document(self.out / 'manual.docx').paragraphs)
                self.assertIn(faqs, docx_text)
                self.assertIn(verified, docx_text)

                with zipfile.ZipFile(self.out / 'manual-markdown.zip') as archive:
                    markdown = archive.read('README.md').decode('utf-8')
                self.assertIn(faqs, markdown)
                self.assertIn(verified, markdown)


if __name__=='__main__': unittest.main()
